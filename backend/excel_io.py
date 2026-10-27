"""Excel 导入导出公共工具：xlsx 读写、单元格解析、HTTP 响应。"""
from __future__ import annotations

from io import BytesIO
from typing import Any, Iterable
from urllib.parse import quote

from django.http import HttpResponse
from rest_framework.decorators import action
from rest_framework.exceptions import ValidationError
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.response import Response

def _openpyxl():
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise ImportError('请安装 openpyxl：pip install openpyxl') from exc
    return Workbook, load_workbook, Alignment, Font, PatternFill, get_column_letter


def cell_str(value: Any) -> str:
    if value is None:
        return ''
    text = str(value).strip()
    if text.lower() in ('none', 'null', 'nan'):
        return ''
    return text


def import_code(value: Any, *, max_length: int = 32, label: str = '编号') -> str:
    """表格中的编号原样使用。Excel 数字单元格的 ``123.0`` 收成 ``123``。"""
    text = cell_str(value)
    if text.endswith('.0') and text[:-2].isdigit():
        text = text[:-2]
    if text and len(text) > max_length:
        raise ValueError(f'{label}不能超过 {max_length} 个字符')
    return text


def as_bool(value: Any, default: bool | None = None) -> bool | None:
    text = cell_str(value).lower()
    if not text:
        return default
    if text in ('1', 'true', 'yes', 'y', '是', '启用', '开', 'on'):
        return True
    if text in ('0', 'false', 'no', 'n', '否', '停用', '关', 'off'):
        return False
    return default


def as_int(value: Any, default: int | None = None) -> int | None:
    text = cell_str(value)
    if not text:
        return default
    try:
        return int(float(text))
    except (TypeError, ValueError):
        return default


def as_decimal(value: Any, default=None):
    from decimal import Decimal, InvalidOperation

    text = cell_str(value)
    if not text:
        return default
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return default


def split_list(value: Any) -> list[str]:
    text = cell_str(value)
    if not text:
        return []
    parts = []
    for raw in text.replace('；', ',').replace(';', ',').replace('/', ',').replace('|', ',').split(','):
        item = raw.strip()
        if item and item not in parts:
            parts.append(item)
    return parts


def map_choice(value: Any, aliases: dict[str, str], default: str = '') -> str:
    text = cell_str(value)
    if not text:
        return default
    key = text.lower()
    if key in aliases:
        return aliases[key]
    if text in aliases:
        return aliases[text]
    return text if text in set(aliases.values()) else default or text


def _autosize(ws, get_column_letter, max_width: int = 36) -> None:
    for idx, col in enumerate(ws.columns, start=1):
        width = 10
        for cell in col:
            width = max(width, min(max_width, len(cell_str(cell.value)) + 2))
        ws.column_dimensions[get_column_letter(idx)].width = width


def write_workbook(
    *,
    filename: str,
    sheets: list[dict[str, Any]],
) -> HttpResponse:
    """sheets: [{name, headers, rows, hints?, example?}]"""
    Workbook, _, Alignment, Font, PatternFill, get_column_letter = _openpyxl()
    header_fill = PatternFill('solid', fgColor='1F4E79')
    header_font = Font(color='FFFFFF', bold=True)
    hint_fill = PatternFill('solid', fgColor='FFF2CC')
    example_fill = PatternFill('solid', fgColor='E2EFDA')
    wb = Workbook()
    first = True
    for spec in sheets:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = spec['name'][:31]
        hints = spec.get('hints') or []
        if hints:
            ws.append(hints if isinstance(hints[0], list) else [hints[0]])
            for cell in ws[1]:
                cell.fill = hint_fill
                cell.alignment = Alignment(wrap_text=True, vertical='top')
            ws.row_dimensions[1].height = 48
            ws.append([])
        headers = spec['headers']
        ws.append(headers)
        header_row = ws.max_row
        for cell in ws[header_row]:
            cell.fill = header_fill
            cell.font = header_font
        example = spec.get('example')
        if example:
            ws.append(example)
            for cell in ws[ws.max_row]:
                cell.fill = example_fill
        for row in spec.get('rows') or []:
            ws.append([row.get(h, '') if isinstance(row, dict) else row[i] for i, h in enumerate(headers)] if not isinstance(row, dict) else [row.get(h, '') for h in headers])
        ws.freeze_panes = ws.cell(header_row + 1, 1).coordinate
        ws.auto_filter.ref = f'A{header_row}:{get_column_letter(len(headers))}{ws.max_row}'
        _autosize(ws, get_column_letter)
    return workbook_response(wb, filename)


def workbook_response(wb, filename: str) -> HttpResponse:
    buf = BytesIO()
    wb.save(buf)
    payload = buf.getvalue()
    response = HttpResponse(
        payload,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    ascii_name = filename.encode('ascii', 'ignore').decode() or 'export.xlsx'
    response['Content-Disposition'] = (
        f'attachment; filename="{ascii_name}"; filename*=UTF-8\'\'{quote(filename)}'
    )
    response['Content-Length'] = str(len(payload))
    return response


def load_upload(request) -> Any:
    upload = request.FILES.get('file') or request.FILES.get('excel')
    if not upload:
        raise ValidationError({'file': '请上传 Excel 文件（.xlsx）'})
    name = (upload.name or '').lower()
    if not name.endswith('.xlsx'):
        raise ValidationError({'file': '仅支持 .xlsx 文件，请先下载模板填写'})
    try:
        _, load_workbook, *_ = _openpyxl()
        return load_workbook(filename=BytesIO(upload.read()), data_only=True)
    except ImportError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ValidationError({'file': f'无法读取 Excel：{exc}'}) from exc


def sheet_rows(wb, names: Iterable[str] | None = None) -> list[dict[str, str]]:
    """读取指定工作表（或第一个非「填写说明」表）为字典列表。"""
    wanted = {n.strip() for n in (names or []) if n}
    ws = None
    for name in wb.sheetnames:
        if wanted and name in wanted:
            ws = wb[name]
            break
        if not wanted and name not in ('填写说明', '说明'):
            ws = wb[name]
            break
    if ws is None:
        ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header_idx = None
    headers: list[str] = []
    for i, row in enumerate(rows):
        values = [cell_str(v) for v in (row or ())]
        if any(values) and not (len(values) == 1 and '填写' in values[0]):
            header_idx = i
            headers = values
            break
    if header_idx is None:
        return []
    out = []
    for row in rows[header_idx + 1:]:
        values = [cell_str(v) for v in (row or ())]
        if not any(values):
            continue
        item = {}
        for idx, key in enumerate(headers):
            if key:
                item[key] = values[idx] if idx < len(values) else ''
        if any(item.values()):
            out.append(item)
    return out


def import_result(created=0, updated=0, skipped=0, errors=None, extra=None) -> dict:
    data = {
        'created': created,
        'updated': updated,
        'skipped': skipped,
        'errors': errors or [],
    }
    if extra:
        data.update(extra)
    return data


class ExcelIOMixin:
    """ViewSet：GET export / GET import-template / POST import。"""

    excel_export_fn = None
    excel_import_fn = None

    def _excel_fn(self, attr: str):
        raw = type(self).__dict__.get(attr)
        if isinstance(raw, staticmethod):
            return raw.__func__
        if callable(raw):
            return raw
        fn = getattr(self, attr, None)
        return fn.__func__ if isinstance(fn, staticmethod) else fn

    def _excel_filtered_qs(self):
        return self.filter_queryset(self.get_queryset())

    def _run_excel_import(self, wb, request):
        import inspect
        fn = self._excel_fn('excel_import_fn')
        if 'user' in inspect.signature(fn).parameters:
            return fn(wb, user=request.user)
        return fn(wb)

    @action(detail=False, methods=['get'], url_path='export')
    def export(self, request):
        return self._excel_fn('excel_export_fn')(self._excel_filtered_qs())

    @action(detail=False, methods=['get'], url_path='import-template')
    def import_template(self, request):
        return self._excel_fn('excel_export_fn')(self.get_queryset().none())

    @action(
        detail=False, methods=['post'], url_path='import',
        parser_classes=[MultiPartParser, FormParser],
    )
    def import_excel(self, request):
        result = self._run_excel_import(load_upload(request), request)
        return Response(result)
