"""产品 / 模版 Excel 导入导出。"""
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from product_manager.models import Product, ProductTemplate, ProductUpgradeRule, TemplatePricing
from resource_manager.models import City, Country, PopPoint, Region


def _xlsx(sheets):
    wb = Workbook()
    first = True
    for name, headers, rows in sheets:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name
        ws.append(headers)
        for row in rows:
            ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = 'data.xlsx'
    return buf


class ProductExcelTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser('excel-prod', 'excel-prod@example.test', 'password')
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        region = Region.objects.create(code='xlp', name='Excel产品区')
        country = Country.objects.create(region=region, name='Excel产品国')
        city = City.objects.create(country=country, name='Excel产品市')
        self.pop = PopPoint.objects.create(name='Excel产品POP', city=city)
        self.product = Product.objects.create(
            name='导出用VPS',
            status=Product.Status.ONLINE,
            pop_point=self.pop,
            cpu_cores=2,
            ram_mb=2048,
            disk_gb=40,
            monthly_price=18,
        )

    def test_export_contains_product_and_headers(self):
        res = self.client.get('/api/product/products/export/')
        self.assertEqual(res.status_code, 200)
        self.assertIn('spreadsheetml', res['Content-Type'])
        wb = load_workbook(BytesIO(res.content))
        self.assertIn('产品', wb.sheetnames)
        rows = list(wb['产品'].iter_rows(values_only=True))
        headers = [str(v) for v in rows[0]]
        self.assertIn('产品编号', headers)
        self.assertIn('POP编号', headers)
        self.assertIn('大区', headers)
        self.assertIn('城市', headers)
        self.assertIn('阶梯定价', wb.sheetnames)
        names = [row[headers.index('产品名称')] for row in rows[1:]]
        self.assertIn('导出用VPS', names)
        loc_row = next(row for row in rows[1:] if row[headers.index('产品名称')] == '导出用VPS')
        self.assertEqual(loc_row[headers.index('大区')], 'Excel产品区')
        self.assertEqual(loc_row[headers.index('城市')], 'Excel产品市')

    def test_import_creates_and_updates(self):
        buf = _xlsx([('产品', [
            '产品编号', '产品名称', '类别', '状态', 'POP编号',
            'CPU核心', '内存MB', '硬盘GB', '网卡数', '带宽Mbps', 'IP数量',
            '月租价', '币种',
        ], [
            ['', '导入新建VPS', 'VPS 主机', '草稿', self.pop.code,
             1, 1024, 20, 1, 10, 1, 9, 'USD'],
        ])])
        res = self.client.post('/api/product/products/import/', {'file': buf}, format='multipart')
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(res.data['created'], 1)
        created = Product.objects.get(name='导入新建VPS')
        self.assertEqual(created.pop_point_id, self.pop.id)
        self.assertEqual(created.status, Product.Status.DRAFT)

        buf = _xlsx([('产品', [
            '产品编号', '产品名称', '类别', '状态', 'POP编号', '月租价',
        ], [
            [created.product_no, '导入更新VPS', 'VPS', '已上线', self.pop.code, 33],
        ])])
        res = self.client.post('/api/product/products/import/', {'file': buf}, format='multipart')
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(res.data['updated'], 1)
        created.refresh_from_db()
        self.assertEqual(created.name, '导入更新VPS')
        self.assertEqual(created.status, Product.Status.ONLINE)
        self.assertEqual(float(created.monthly_price), 33)

    def test_sheet_product_no_is_kept_and_reimport_updates(self):
        buf = _xlsx([('产品', [
            '产品编号', '产品名称', '类别', '状态', 'POP编号', '月租价',
        ], [
            ['SKU-VPS-01', '编号新建', 'VPS', '草稿', self.pop.code, 8],
        ])])
        res = self.client.post('/api/product/products/import/', {'file': buf}, format='multipart')
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(res.data['created'], 1)
        created = Product.objects.get(product_no='SKU-VPS-01')
        self.assertEqual(created.name, '编号新建')

        buf = _xlsx([('产品', [
            '产品编号', '产品名称', '类别', '状态', 'POP编号', '月租价',
        ], [
            ['SKU-VPS-01', '编号覆盖', 'VPS', '已上线', self.pop.code, 18],
        ])])
        res = self.client.post('/api/product/products/import/', {'file': buf}, format='multipart')
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(res.data['updated'], 1)
        self.assertEqual(res.data['created'], 0)
        self.assertEqual(Product.objects.filter(product_no='SKU-VPS-01').count(), 1)
        created.refresh_from_db()
        self.assertEqual(created.name, '编号覆盖')
        self.assertEqual(float(created.monthly_price), 18)

    def test_import_rejects_non_xlsx(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        f = SimpleUploadedFile('data.csv', b'a,b\n1,2', content_type='text/csv')
        res = self.client.post('/api/product/products/import/', {'file': f}, format='multipart')
        self.assertEqual(res.status_code, 400)

    def test_template_roundtrip_pricing(self):
        tpl = ProductTemplate.objects.create(name='Excel模版1C', category='vps', cpu_cores=1, ram_mb=1024)
        TemplatePricing.objects.create(template=tpl, pop_point=self.pop, monthly_price=12, currency='USD')
        res = self.client.get('/api/product/templates/export/')
        self.assertEqual(res.status_code, 200)
        wb = load_workbook(BytesIO(res.content))
        self.assertIn('模版', wb.sheetnames)
        self.assertIn('定价', wb.sheetnames)

        buf = _xlsx([
            ('模版', ['模版编号', '模版名称', '类别', 'CPU核心', '内存MB', '硬盘GB', '启用'], [
                ['', '导入模版2C', 'VPS', 2, 2048, 40, '是'],
            ]),
            ('定价', ['模版编号', 'POP编号', '月租价', '信誉值', '安装费', '币种'], [
                ['导入模版2C', self.pop.code, 25, 0, 0, 'USD'],
            ]),
        ])
        res = self.client.post('/api/product/templates/import/', {'file': buf}, format='multipart')
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(res.data['created'], 1)
        obj = ProductTemplate.objects.get(name='导入模版2C')
        self.assertGreaterEqual(res.data['pricing_created'] + res.data['pricing_updated'], 1)
        self.assertTrue(TemplatePricing.objects.filter(template=obj, pop_point=self.pop, monthly_price=25).exists())

    def test_import_upgrade_rules_and_location_pop(self):
        buf = _xlsx([
            ('产品', [
                '产品名称', '类别', '状态', 'POP名称', '大区', '国家', '城市', '月租价',
            ], [
                ['阶梯VPS', 'VPS', '草稿', self.pop.name, 'Excel产品区', 'Excel产品国', 'Excel产品市', 20],
            ]),
            ('阶梯定价', [
                '产品名称', '资源类型', '起始数量', '结束数量', '每单位月租', '每单位信誉值', '启用',
            ], [
                ['阶梯VPS', 'CPU 核', 1, 4, 2.5, 0, '是'],
                ['阶梯VPS', '带宽(Mbps)', 10, '', 0.2, 0, '是'],
            ]),
        ])
        res = self.client.post('/api/product/products/import/', {'file': buf}, format='multipart')
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(res.data['created'], 1)
        self.assertEqual(res.data['rules_created'], 2)
        product = Product.objects.get(name='阶梯VPS')
        self.assertEqual(product.pop_point_id, self.pop.id)
        cpu = ProductUpgradeRule.objects.get(product=product, resource_kind='cpu')
        self.assertEqual(cpu.min_qty, 1)
        self.assertEqual(cpu.max_qty, 4)
        self.assertEqual(float(cpu.unit_price), 2.5)
        bw = ProductUpgradeRule.objects.get(product=product, resource_kind='bandwidth')
        self.assertIsNone(bw.max_qty)

    def test_edge_category_survives_template_and_product_reimport(self):
        tpl = ProductTemplate.objects.create(
            name='边缘往返模版', category='edge', accel_type='gre,wireguard',
            cpu_cores=0, ram_mb=0, disk_gb=0, nic_count=0, bandwidth_mbps=20, ip_count=1,
        )
        product = Product.objects.create(
            name='边缘往返产品', category='edge', status=Product.Status.ONLINE,
            pop_point=self.pop, template=tpl, accel_type='gre,wireguard',
            cpu_cores=0, ram_mb=0, disk_gb=0, bandwidth_mbps=20, ip_count=1,
            monthly_price=8,
        )
        tpl_res = self.client.get('/api/product/templates/export/')
        self.assertEqual(tpl_res.status_code, 200)
        tpl_wb = load_workbook(BytesIO(tpl_res.content))
        headers = [str(v) for v in next(tpl_wb['模版'].iter_rows(values_only=True))]
        row = next(
            r for r in tpl_wb['模版'].iter_rows(values_only=True)
            if r[headers.index('模版名称')] == '边缘往返模版'
        )
        self.assertEqual(row[headers.index('类别')], '边缘加速')
        self.assertIn('内存GB', headers)

        buf = _xlsx([
            ('模版', ['模版编号', '模版名称', '类别', '带宽Mbps', 'IP数量', '加速类型', '启用'], [
                [tpl.template_no, '边缘往返模版', '边缘加速', 30, 2, 'gre,wireguard', '是'],
                [tpl.template_no, '边缘类别留空', '', 30, 2, 'gre', '是'],
            ]),
        ])
        res = self.client.post('/api/product/templates/import/', {'file': buf}, format='multipart')
        self.assertEqual(res.status_code, 200, res.data)
        tpl.refresh_from_db()
        self.assertEqual(tpl.category, 'edge')
        self.assertEqual(tpl.name, '边缘类别留空')
        self.assertIn('gre', tpl.accel_type)

        buf = _xlsx([
            ('产品', ['产品编号', '产品名称', '类别', 'POP编号', '内存GB', '月租价'], [
                [product.product_no, '边缘往返产品', '边缘加速', self.pop.code, 4, 12],
            ]),
        ])
        res = self.client.post('/api/product/products/import/', {'file': buf}, format='multipart')
        self.assertEqual(res.status_code, 200, res.data)
        product.refresh_from_db()
        self.assertEqual(product.category, 'edge')
        self.assertEqual(product.ram_mb, 0)

        vps = Product.objects.create(
            name='内存GB产品', category='vps', pop_point=self.pop, ram_mb=1024, monthly_price=1,
        )
        buf = _xlsx([
            ('产品', ['产品编号', '产品名称', '类别', 'POP编号', '内存GB'], [
                [vps.product_no, '内存GB产品', 'VPS', self.pop.code, 2],
            ]),
        ])
        res = self.client.post('/api/product/products/import/', {'file': buf}, format='multipart')
        self.assertEqual(res.status_code, 200, res.data)
        vps.refresh_from_db()
        self.assertEqual(vps.category, 'vps')
        self.assertEqual(vps.ram_mb, 2048)
