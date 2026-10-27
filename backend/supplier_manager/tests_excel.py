"""供应商列表 Excel 导入导出。"""
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from supplier_manager.models import Supplier


def _xlsx(sheets):
    wb = Workbook()
    first = True
    for name, headers, rows in sheets:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name
        ws.append(headers)
        for row in rows:
            ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = 'data.xlsx'
    return buf


class SupplierExcelTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser('excel-sup', 'excel-sup@example.test', 'password')
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.supplier = Supplier.objects.create(
            name='导出供应商',
            supplier_type='cloud',
            status='active',
            contact_name='李四',
            settlement_currency='USD',
        )

    def test_export_contains_supplier(self):
        res = self.client.get('/api/supplier/suppliers/export/')
        self.assertEqual(res.status_code, 200, res.content[:200])
        self.assertIn('spreadsheetml', res['Content-Type'])
        wb = load_workbook(BytesIO(res.content))
        self.assertIn('供应商', wb.sheetnames)
        rows = list(wb['供应商'].iter_rows(values_only=True))
        headers = [str(v) for v in rows[1 if rows[0] and rows[0][0] != '供应商编号' else 0]]
        # 示例行在表头后，数据跟在后面
        header_row = next(i for i, row in enumerate(rows) if row and row[0] == '供应商编号')
        headers = [str(v) for v in rows[header_row]]
        names = [row[headers.index('供应商名称')] for row in rows[header_row + 1:]]
        self.assertIn('导出供应商', names)

    def test_import_creates_and_updates(self):
        buf = _xlsx([('供应商', [
            '供应商编号', '供应商名称', '类型', '状态', '联系人', '结算币种', '信誉评分',
        ], [
            ['', '导入新建供应商', '云厂商', '合作中', '王五', 'CNY', 80],
        ])])
        res = self.client.post('/api/supplier/suppliers/import/', {'file': buf}, format='multipart')
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(res.data['created'], 1)
        created = Supplier.objects.get(name='导入新建供应商')
        self.assertEqual(created.supplier_type, 'cloud')
        self.assertEqual(created.status, 'active')
        self.assertEqual(created.rating, 80)
        self.assertTrue(created.supplier_no)

        buf = _xlsx([('供应商', [
            '供应商编号', '供应商名称', '类型', '状态', '联系人',
        ], [
            [created.supplier_no, '导入更新供应商', '代理商', '暂停合作', '赵六'],
        ])])
        res = self.client.post('/api/supplier/suppliers/import/', {'file': buf}, format='multipart')
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(res.data['updated'], 1)
        created.refresh_from_db()
        self.assertEqual(created.name, '导入更新供应商')
        self.assertEqual(created.supplier_type, 'agent')
        self.assertEqual(created.status, 'suspended')
        self.assertEqual(created.contact_name, '赵六')

    def test_sheet_supplier_no_is_kept_and_reimport_updates(self):
        buf = _xlsx([('供应商', [
            '供应商编号', '供应商名称', '类型', '状态',
        ], [
            ['SUP-KEEP-01', '编号供应商', '云厂商', '合作中'],
        ])])
        res = self.client.post('/api/supplier/suppliers/import/', {'file': buf}, format='multipart')
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(res.data['created'], 1)
        created = Supplier.objects.get(supplier_no='SUP-KEEP-01')

        buf = _xlsx([('供应商', [
            '供应商编号', '供应商名称', '类型', '状态',
        ], [
            ['SUP-KEEP-01', '编号供应商已改', '代理商', '暂停合作'],
        ])])
        res = self.client.post('/api/supplier/suppliers/import/', {'file': buf}, format='multipart')
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(res.data['updated'], 1)
        self.assertEqual(res.data['created'], 0)
        self.assertEqual(Supplier.objects.filter(supplier_no='SUP-KEEP-01').count(), 1)
        created.refresh_from_db()
        self.assertEqual(created.name, '编号供应商已改')
        self.assertEqual(created.supplier_type, 'agent')

    def test_template_skips_example_row(self):
        res = self.client.get('/api/supplier/suppliers/import-template/')
        self.assertEqual(res.status_code, 200)
        wb = load_workbook(BytesIO(res.content))
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        buf.name = 'template.xlsx'
        res = self.client.post('/api/supplier/suppliers/import/', {'file': buf}, format='multipart')
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(res.data['created'], 0)
        self.assertFalse(Supplier.objects.filter(name__startswith='示例').exists())
